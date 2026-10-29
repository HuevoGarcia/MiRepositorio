import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from monaco import Registro, ValidationError, generar_sql, guardar_sql, leer_excel, main, precio


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "tarifas.xlsx"

    def excel(self, rows, headers=None):
        book = Workbook()
        book.active.append(["FORMULARIO NUEVAS TARIFAS LDI"])
        book.active.merge_cells("A1:D1")
        book.active.append(headers or ["Pais", "Destino", "Características", "Precio (USD)"])
        for row in rows:
            book.active.append(row)
        book.save(self.path)
        book.close()
        return self.path

    def test_expansion_y_ceros_iniciales(self):
        path = self.excel([["AFGHANISTAN", "AWCC", "09370, 9371", "0,16330"]])
        result = leer_excel(path)
        self.assertEqual([r.values() for r in result.registros], [
            ("AFGHANISTAN", "AWCC", "09370", "0.1633"),
            ("AFGHANISTAN", "AWCC", "9371", "0.1633"),
        ])

    def test_precio_numerico_y_codigo_entero(self):
        result = leer_excel(self.excel([["Malta", "Malta", 356, 0.18]]))
        self.assertEqual(result.registros[0].values(), ("Malta", "Malta", "356", "0.18"))

    def test_cada_campo_es_obligatorio(self):
        for index in range(4):
            with self.subTest(campo=index):
                row = ["Malta", "Malta", "356", "0,18"]
                row[index] = None
                with self.assertRaisesRegex(ValidationError, "Fila 3"):
                    leer_excel(self.excel([row]))

    def test_segmento_vacio_rechaza_archivo(self):
        for code in ("356,", ",356", "356,,357"):
            with self.subTest(code=code), self.assertRaises(ValidationError):
                leer_excel(self.excel([["Malta", "Malta", code, "0,18"]]))

    def test_duplicados_identicos_se_omiten(self):
        row = ["Malta", "Malta", "356", "0,18"]
        result = leer_excel(self.excel([row, row]))
        self.assertEqual(len(result.registros), 1)
        self.assertEqual(result.repetidos, 1)

    def test_clave_repetida_precio_distinto_rechaza_archivo(self):
        with self.assertRaisesRegex(ValidationError, r"Fila 4:.*precios diferentes.*fila 3"):
            leer_excel(self.excel([["Malta", "Malta", "356", "0,18"], ["Malta", "Malta", "356", "0,20"]]))

    def test_misma_caracteristica_otro_destino_es_otra_clave(self):
        result = leer_excel(self.excel([["Malta", "Fijo", "356", "0,18"], ["Malta", "Movil", "356", "0,20"]]))
        self.assertEqual(len(result.registros), 2)

    def test_formulas_y_errores_rechazados(self):
        for value in ("=1/2", "#DIV/0!"):
            with self.subTest(value=value), self.assertRaisesRegex(ValidationError, "fórmula o error"):
                leer_excel(self.excel([["Malta", "Malta", "356", value]]))

    def test_limites_varchar(self):
        for country in ("a" * 256, "日本"):
            with self.subTest(country=country), self.assertRaises(ValidationError):
                leer_excel(self.excel([[country, "Malta", "356", "0,18"]]))

    def test_hoja_y_encabezados(self):
        path = self.excel([["Malta", "Malta", "356", "0,18"]], ["Pais", "Destino", "Caracteristicas", "Incorrecto"])
        with self.assertRaisesRegex(ValidationError, "fila 2"):
            leer_excel(path)
        with self.assertRaisesRegex(ValidationError, "No existe la hoja"):
            leer_excel(path, "NoExiste")

    def test_fila_1_se_ignora_aunque_contenga_otros_encabezados(self):
        path = self.excel([["Malta", "Malta", "356", "0,18"]])
        from openpyxl import load_workbook
        book = load_workbook(path)
        book.active.unmerge_cells("A1:D1")
        for index, value in enumerate(["Precio (USD)", "Caracteristicas", "Destino", "Pais"], start=1):
            book.active.cell(1, index, value)
        book.save(path)
        book.close()
        result = leer_excel(path)
        self.assertEqual(result.filas, 1)
        self.assertEqual(result.registros[0].values(), ("Malta", "Malta", "356", "0.18"))

    def test_titulo_sin_encabezados_rechazado(self):
        book = Workbook()
        book.active.append(["FORMULARIO NUEVAS TARIFAS LDI"])
        book.save(self.path)
        book.close()
        with self.assertRaisesRegex(ValidationError, "fila 2"):
            leer_excel(self.path)

    def test_filas_blancas_y_archivo_sin_registros(self):
        with self.assertRaisesRegex(ValidationError, "no contiene registros"):
            leer_excel(self.excel([[None, None, None, None]]))
        result = leer_excel(self.excel([[None] * 4, ["Malta", "Malta", "356", "0,18"]]))
        self.assertEqual(result.filas, 1)

    def test_precios_invalidos_y_precision(self):
        for value in ("-1", "1.000,25", "NaN", "Infinity", True, "USD 1", " "):
            with self.subTest(value=value), self.assertRaises(ValidationError):
                precio(value)
        self.assertEqual(precio("0,12345678901234567890"), "0.1234567890123456789")
        self.assertEqual(precio(0), "0")
        self.assertEqual(precio(0.00001), "0.00001")

    def test_precio_normalizado_con_punto_para_sql(self):
        for value in ("0,16330", "0.16330", 0.1633):
            with self.subTest(value=value):
                planilla = leer_excel(self.excel([["Malta", "Malta", "356", value]]))
                sql = generar_sql(planilla.registros)
                self.assertIn("(N'Malta', N'Malta', N'356', N'0.1633')", sql)
                self.assertNotIn("N'0,16330'", sql)

    def test_cli_genera_sql_y_no_sobrescribe_archivos(self):
        path = self.excel([["Malta", "Malta", "356", "0,18"]])
        output = Path(self.temp.name) / "tarifas.sql"
        self.assertEqual(main(["--archivo", str(path), "--salida", str(output)]), 0)
        content = output.read_bytes()
        self.assertTrue(content.startswith(b"\xef\xbb\xbf"))
        self.assertIn("N'0.18'", content.decode("utf-8-sig"))
        output.write_text("SQL revisado por usuario")
        self.assertEqual(main(["--archivo", str(path), "--salida", str(output)]), 1)
        self.assertEqual(output.read_text(), "SQL revisado por usuario")

    def test_cli_excel_invalido_no_genera_sql(self):
        path = self.excel([["Malta", "Malta", "356", None]])
        output = Path(self.temp.name) / "tarifas.sql"
        self.assertEqual(main(["--archivo", str(path), "--salida", str(output)]), 1)
        self.assertFalse(output.exists())

    def test_solo_validar_no_genera_sql(self):
        path = self.excel([["Malta", "Malta", "356", "0,18"]])
        output = Path(self.temp.name) / "tarifas.sql"
        self.assertEqual(main(["--archivo", str(path), "--salida", str(output), "--validar"]), 0)
        self.assertFalse(output.exists())


class SQLTests(unittest.TestCase):
    def test_literales_no_permiten_escapar_del_valor(self):
        sql = generar_sql([Registro("Côte d'Ivoire", "x'); DROP TABLE dbo.PreciosMonaco;--", "001", "0.18")])
        self.assertIn("N'Côte d''Ivoire'", sql)
        self.assertIn("N'x''); DROP TABLE dbo.PreciosMonaco;--'", sql)
        self.assertIn("N'001'", sql)

    def test_planilla_grande_respeta_limite_insert_values(self):
        records = [Registro("Malta", "Malta", str(i), "0.18") for i in range(6018)]
        sql = generar_sql(records)
        batches = sql.split("INSERT INTO #MonacoCarga (Pais, Destino, Caracteristicas, Precio) VALUES\n")[1:]
        self.assertEqual(len(batches), 13)
        sizes = [batch.split(";", 1)[0].count("(N'Malta'") for batch in batches]
        self.assertEqual(sum(sizes), 6018)
        self.assertLessEqual(max(sizes), 1000)
        self.assertIn("N'6017'", sql)

    def test_destino_clave_y_confirmacion_en_sql(self):
        sql = generar_sql([Registro("Malta", "Malta", "356", "0.18")])
        self.assertIn("USE [Auxiliar];", sql)
        self.assertIn("DECLARE @Aplicar bit = 0;", sql)
        branch = sql.split("IF @Aplicar = 1", 1)[1]
        self.assertIn("UPDATE t SET Precio = s.Precio", branch)
        self.assertIn("INSERT INTO dbo.PreciosMonaco", branch)
        self.assertIn("WHERE NOT EXISTS", branch)
        self.assertIn("t.Pais = s.Pais AND t.Destino = s.Destino AND t.Caracteristicas = s.Caracteristicas", branch)
        self.assertIn("ELSE\n        ROLLBACK TRANSACTION;", branch)
        self.assertIn("IF XACT_STATE() <> 0 ROLLBACK TRANSACTION;", branch)
        self.assertIn("WITH (TABLOCKX, HOLDLOCK)", sql)
        self.assertEqual(sql.count("HAVING COUNT(*) > 1"), 2)

    def test_sin_registros_rechazado(self):
        with self.assertRaises(ValidationError):
            generar_sql([])


if __name__ == "__main__":
    unittest.main()
