"""Importación de tarifas Monaco desde Excel a SQL Server (Python 3.14)."""

import argparse
import configparser
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
import logging
import math
from pathlib import Path
import re
import sys
import unicodedata

BASE_DIR = Path(__file__).resolve().parent


class ValidationError(ValueError):
    pass


@dataclass(frozen=True)
class Registro:
    pais: str
    destino: str
    caracteristicas: str
    precio: str

    def values(self):
        return self.pais, self.destino, self.caracteristicas, self.precio


@dataclass
class Planilla:
    hoja: str
    filas: int
    repetidos: int
    registros: list[Registro]


def texto(value, campo):
    if value is None or isinstance(value, bool):
        raise ValidationError(f"{campo}: valor vacío o inválido.")
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValidationError(f"{campo}: número inválido.")
        # Excel solo conserva 15 dígitos significativos en celdas numéricas.
        value = int(value) if value.is_integer() else value
    result = str(value).strip()
    if not result or any(ord(c) < 32 for c in result):
        raise ValidationError(f"{campo}: valor vacío o caracteres de control.")
    try:
        encoded = result.encode("cp1252")
    except UnicodeEncodeError as exc:
        raise ValidationError(f"{campo}: contiene caracteres incompatibles con varchar de la tabla.") from exc
    if len(encoded) > 255:
        raise ValidationError(f"{campo}: supera los 255 caracteres permitidos.")
    return result


def precio(value):
    # str(float) puede usar exponente para tarifas pequeñas (p. ej. 0.00001).
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValidationError("Precio: número inválido.")
        value = format(Decimal(str(value)), "f")
    raw = texto(value, "Precio")
    if not re.fullmatch(r"[0-9]+(?:[.,][0-9]+)?", raw):
        raise ValidationError("Precio: debe ser un número no negativo, sin separadores de miles ni moneda.")
    try:
        number = Decimal(raw.replace(",", "."))
    except InvalidOperation as exc:
        raise ValidationError("Precio: número inválido.") from exc
    result = format(number, "f")
    if "." in result:
        result = result.rstrip("0").rstrip(".")
    return texto(result.replace(".", ","), "Precio")


def encabezado(value):
    raw = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join("".join(c for c in raw if not unicodedata.combining(c)).casefold().split())


def leer_excel(path, hoja=""):
    from openpyxl import load_workbook

    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValidationError("Seleccioná un archivo .xlsx; convertí los .xls antes de cargar.")
    # data_only=False permite rechazar fórmulas sin depender de cachés de Excel.
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if hoja and hoja not in workbook.sheetnames:
            raise ValidationError(f"No existe la hoja '{hoja}'.")
        sheet = workbook[hoja] if hoja else workbook.active
        # La fila 1 es el título; la fila 2 contiene los encabezados.
        rows = sheet.iter_rows(min_row=2)
        first = next(rows, None)
        if not first:
            raise ValidationError("Faltan los encabezados en la fila 2.")
        wanted = {"pais": "Pais", "destino": "Destino", "caracteristicas": "Caracteristicas", "precio (usd)": "Precio"}
        columns = {}
        for index, cell in enumerate(first):
            name = wanted.get(encabezado(cell.value))
            if name:
                if name in columns:
                    raise ValidationError(f"Encabezado repetido: {name}.")
                columns[name] = index
        if len(columns) != 4:
            raise ValidationError("La fila 2 debe contener Pais, Destino, Características y Precio (USD).")
        records = {}
        source_rows = {}
        errors = []
        count = duplicates = 0
        for row_number, cells in enumerate(rows, start=3):
            # Ignorar solo filas totalmente vacías, no registros incompletos.
            if all(c.value is None or (isinstance(c.value, str) and not c.value.strip()) for c in cells):
                continue
            count += 1
            try:
                values = {}
                for name, index in columns.items():
                    cell = cells[index] if index < len(cells) else None
                    if cell is not None and cell.data_type in ("f", "e"):
                        raise ValidationError(f"{name}: fórmula o error de Excel; usar valores.")
                    values[name] = cell.value if cell else None
                pais = texto(values["Pais"], "Pais")
                destino = texto(values["Destino"], "Destino")
                cost = precio(values["Precio"])
                raw = values["Caracteristicas"]
                if raw is None:
                    raise ValidationError("Caracteristicas: valor vacío.")
                parts = raw.split(",") if isinstance(raw, str) else [raw]
                for part in parts:
                    feature = texto(part, "Caracteristicas")
                    key = pais, destino, feature
                    if key in records:
                        if records[key].precio != cost:
                            raise ValidationError(f"Clave repetida con precios diferentes (también en fila {source_rows[key]}).")
                        duplicates += 1
                    else:
                        records[key] = Registro(pais, destino, feature, cost)
                        source_rows[key] = row_number
            except ValidationError as exc:
                errors.append(f"Fila {row_number}: {exc}")
        if errors:
            raise ValidationError("Planilla rechazada. No se cargó ningún dato.\n" + "\n".join(errors))
        if not records:
            raise ValidationError("La planilla no contiene registros para cargar.")
        return Planilla(sheet.title, count, duplicates, list(records.values()))
    finally:
        workbook.close()


def configuracion(path):
    config = configparser.ConfigParser(interpolation=None)
    if not config.read(path, encoding="utf-8-sig"):
        raise ValidationError(f"No se encontró la configuración: {path}")
    for key in ("servidor", "base", "driver"):
        if not config.get("sqlserver", key, fallback="").strip():
            raise ValidationError(f"Falta sqlserver.{key} en la configuración.")
    if config.get("sqlserver", "base") != "Auxiliar":
        raise ValidationError("La base configurada debe ser Auxiliar.")
    return config


def conectar(config):
    import pyodbc

    section = config["sqlserver"]
    driver = section["driver"]
    if driver not in pyodbc.drivers():
        raise ValidationError(f"No está instalado '{driver}'. Instalá Microsoft ODBC Driver 18 for SQL Server (64 bits).")
    def quote(value):
        return "{" + value.replace("}", "}}") + "}"
    encrypt = section.get("encrypt", "yes").lower()
    trust = section.get("trust_server_certificate", "no").lower()
    if encrypt not in ("yes", "mandatory", "strict") or trust not in ("yes", "no"):
        raise ValidationError("Configuración TLS inválida; encrypt debe ser yes, mandatory o strict.")
    connection = (
        f"DRIVER={quote(driver)};SERVER={quote(section['servidor'])};"
        f"DATABASE={quote(section['base'])};Trusted_Connection=yes;"
        f"Encrypt={encrypt};TrustServerCertificate={trust};"
    )
    return pyodbc.connect(connection, autocommit=False, timeout=section.getint("timeout", 30))


MATCH = "t.Pais = s.Pais AND t.Destino = s.Destino AND t.Caracteristicas = s.Caracteristicas"


def cargar(connection, registros, aplicar=False):
    """Resumen coherente y carga atómica; bloquea el destino para evitar carreras.

    La comparación usa la intercalación real de cada columna de SQL Server.
    El bloqueo se libera antes de pedir confirmación (la vista previa hace rollback).
    """
    cursor = connection.cursor()
    try:
        cursor.execute("SET XACT_ABORT ON; SET LOCK_TIMEOUT 30000;")
        # Clonar tipos e intercalaciones, sin cambiar la tabla real.
        cursor.execute("SELECT TOP (0) Pais, Destino, Caracteristicas, Precio INTO #MonacoCarga FROM dbo.PreciosMonaco;")
        cursor.executemany(
            "INSERT INTO #MonacoCarga (Pais, Destino, Caracteristicas, Precio) VALUES (?, ?, ?, ?)",
            [r.values() for r in registros],
        )
        cursor.execute("SELECT TOP (1) 1 FROM #MonacoCarga GROUP BY Pais, Destino, Caracteristicas HAVING COUNT(*) > 1;")
        if cursor.fetchone():
            raise ValidationError("El Excel tiene claves repetidas según la intercalación SQL (por ejemplo, diferencias de mayúsculas). Unificá esas filas.")
        # Se requiere permiso SELECT/INSERT/UPDATE; no se crea ningún índice permanente.
        cursor.execute("SELECT COUNT_BIG(*) FROM dbo.PreciosMonaco WITH (TABLOCKX, HOLDLOCK);")
        cursor.fetchone()
        cursor.execute(f"""
            SELECT TOP (1) 1 FROM dbo.PreciosMonaco t
            JOIN #MonacoCarga s ON {MATCH}
            GROUP BY t.Pais, t.Destino, t.Caracteristicas HAVING COUNT(*) > 1;
        """)
        if cursor.fetchone():
            raise ValidationError("La tabla tiene varias filas para una clave de esta carga. Resolver esos duplicados antes de importar; no se modificó ningún dato.")
        cursor.execute(f"""
            SELECT
                COALESCE(SUM(CASE WHEN t.Pais IS NULL THEN 1 ELSE 0 END), 0),
                COALESCE(SUM(CASE WHEN t.Pais IS NOT NULL AND (t.Precio IS NULL OR t.Precio <> s.Precio) THEN 1 ELSE 0 END), 0),
                COALESCE(SUM(CASE WHEN t.Pais IS NOT NULL AND t.Precio = s.Precio THEN 1 ELSE 0 END), 0)
            FROM #MonacoCarga s LEFT JOIN dbo.PreciosMonaco t ON {MATCH};
        """)
        summary = tuple(cursor.fetchone())
        if aplicar:
            cursor.execute(f"""
                UPDATE t SET Precio = s.Precio FROM dbo.PreciosMonaco t
                JOIN #MonacoCarga s ON {MATCH}
                WHERE t.Precio IS NULL OR t.Precio <> s.Precio;
            """)
            cursor.execute(f"""
                INSERT INTO dbo.PreciosMonaco (Pais, Destino, Caracteristicas, Precio)
                SELECT s.Pais, s.Destino, s.Caracteristicas, s.Precio FROM #MonacoCarga s
                WHERE NOT EXISTS (SELECT 1 FROM dbo.PreciosMonaco t WHERE {MATCH});
            """)
            cursor.execute("DROP TABLE #MonacoCarga;")
            connection.commit()
        else:
            connection.rollback()
        return summary
    except BaseException:
        connection.rollback()
        raise
    finally:
        cursor.close()


def seleccionar_archivo():
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        try:
            return filedialog.askopenfilename(title="Seleccionar planilla Monaco", filetypes=[("Excel", "*.xlsx")])
        finally:
            root.destroy()
    except (ImportError, RuntimeError):
        return input("Ruta del archivo .xlsx: ").strip().strip('"')
    except Exception as exc:
        # Incluye TclError si no existe escritorio en el equipo.
        logging.info("Selector gráfico no disponible: %s", type(exc).__name__)
        return input("Ruta del archivo .xlsx: ").strip().strip('"')


def resumen(result):
    return f"Nuevos: {result[0]} | Precios actualizados: {result[1]} | Sin cambios: {result[2]}"


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archivo", type=Path, help="Planilla .xlsx; si se omite, abre el selector.")
    parser.add_argument("--config", type=Path, default=BASE_DIR / "config.ini")
    parser.add_argument("--validar", action="store_true", help="Validar Excel sin conexión ni cambios en SQL Server.")
    args = parser.parse_args(argv)
    log_dir = BASE_DIR / "logs"
    log_dir.mkdir(exist_ok=True)
    log_path = log_dir / f"monaco_{datetime.now():%Y%m%d_%H%M%S_%f}.log"
    logging.basicConfig(filename=log_path, encoding="utf-8", level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", force=True)
    connection = None
    try:
        path = args.archivo or seleccionar_archivo()
        if not path:
            print("Selección cancelada. No se modificó la base.")
            return 0
        config = configuracion(args.config)
        planilla = leer_excel(path, config.get("excel", "hoja", fallback="").strip())
        message = f"Hoja: {planilla.hoja} | Filas: {planilla.filas} | Registros: {len(planilla.registros)} | Repetidos idénticos omitidos: {planilla.repetidos}"
        print(message)
        logging.info(message)
        if args.validar:
            print("Excel válido. No se conectó ni modificó SQL Server.")
            return 0
        connection = conectar(config)
        print(f"Destino: {config['sqlserver']['servidor']} / Auxiliar.dbo.PreciosMonaco")
        preview = cargar(connection, planilla.registros)
        print("Vista previa: " + resumen(preview))
        logging.info("Vista previa: %s", resumen(preview))
        if input("Escribí SI para confirmar la carga: ").strip().upper() != "SI":
            print("Carga cancelada. No se modificó la base.")
            logging.info("Carga cancelada por el usuario.")
            return 0
        result = cargar(connection, planilla.registros, aplicar=True)
        print("Carga completada: " + resumen(result))
        logging.info("Carga confirmada: %s", resumen(result))
        return 0
    except (KeyboardInterrupt, EOFError):
        print("Operación cancelada.")
        logging.info("Operación interrumpida.")
        return 1
    except Exception as exc:
        # No registrar cadenas de conexión ni credenciales.
        print(f"ERROR: {exc}")
        logging.error("%s: %s", type(exc).__name__, exc)
        return 1
    finally:
        if connection is not None:
            connection.close()
        print(f"Registro de ejecución: {log_path}")


if __name__ == "__main__":
    sys.exit(main())
